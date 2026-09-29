from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
DOCX = OUT / "SIH26036_MeasureSure_Project_Brief.docx"
PDF = OUT / "SIH26036_MeasureSure_Project_Brief.pdf"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_text(cell, text, color="17202B", bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]; r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color); r.font.name = "Aptos"
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.space_before = Pt(16 if level == 1 else 10); p.paragraph_format.space_after = Pt(6); r = p.add_run(text); r.font.name = "Aptos Display"; r.font.color.rgb = RGBColor(28, 60, 98); return p

def build_docx():
    doc = Document(); sec = doc.sections[0]; sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)
    styles = doc.styles; styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10); styles['Normal'].font.color.rgb = RGBColor(77, 91, 108)
    for name in ['Title','Heading 1','Heading 2']:
        styles[name].font.name = 'Aptos Display'
    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = title.add_run('MeasureSure'); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = RGBColor(37,99,235)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run('SIH26036 | Online Verification System for Weighing and Measuring Instruments'); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor(28,49,76)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run('Project brief and implementation concept | 29 September 2026').italic = True
    doc.add_paragraph('MeasureSure is a secure, role-based digital portal that helps citizens, businesses, and Legal Metrology inspectors register, verify, monitor, and renew weighing and measuring instruments. The product is designed around a simple public promise: every measurement should be traceable to a current, trusted verification record.')
    add_heading(doc, '1. Problem statement')
    doc.add_paragraph('Verification records are often spread across paper certificates, local registers, and disconnected spreadsheets. This makes it difficult to confirm whether an instrument is valid at the point of use, creates avoidable work for inspectors, and increases the risk of expired or tampered certificates being accepted.')
    add_heading(doc, '2. Proposed solution')
    doc.add_paragraph('MeasureSure creates one digital record per instrument. Each record carries a unique certificate ID, owner and location details, instrument type, verification history, renewal date, and a digitally signed status. The public search experience is intentionally lightweight: a user can search by certificate number, serial number, or owner name and immediately see whether the instrument is valid, due soon, or pending review.')
    add_heading(doc, '3. Primary users and journeys')
    rows = [['User', 'Primary journey', 'Outcome'], ['Citizen / buyer', 'Search certificate before purchase or transaction', 'Quick confidence in the reading'], ['Business owner', 'Register instrument and track renewal due dates', 'Fewer missed renewals'], ['Inspector', 'Review queue, verify on site, issue certificate', 'Faster, auditable field work'], ['Department admin', 'Monitor district activity and exceptions', 'Clear operational visibility']]
    table = doc.add_table(rows=1, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = 'Table Grid'
    for i, h in enumerate(rows[0]): set_cell_text(table.rows[0].cells[i], h, 'FFFFFF', True, 9); shade(table.rows[0].cells[i], '1C4775')
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, val in enumerate(row): set_cell_text(cells[i], val, size=9)
    add_heading(doc, '4. Product scope')
    for text in ['Public verification search with certificate, serial, and owner lookup.', 'Instrument registration with owner, location, type, and supporting evidence.', 'Inspector workspace with inspection queue, verification actions, and certificate issuance.', 'Certificate repository with status, audit trail, renewal dates, and downloadable records.', 'Dashboard metrics for registered, verified, due-soon, and pending instruments.', 'Tamper-evident digital signatures and role-based access control.']:
        p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3); p.add_run(text)
    add_heading(doc, '5. Prototype interface')
    doc.add_paragraph('The delivered website prototype uses a calm government-service visual language: a clear left navigation, a public verification panel, a compact operational dashboard, and a recent-records table. The main interaction is live in the browser: type a certificate or owner query to filter the record set, use the navigation to change workspace context, or trigger registration and record actions.')
    add_heading(doc, '6. Suggested technical architecture')
    tech = [['Layer', 'Recommended implementation'], ['Frontend', 'Next.js App Router, responsive React components, accessible form controls'], ['API', 'REST or typed RPC endpoints for registration, lookup, verification, and renewal'], ['Data', 'PostgreSQL with instrument, owner, certificate, inspection, and audit tables'], ['Security', 'Role-based auth, signed certificate payloads, encrypted storage, immutable audit events'], ['Deployment', 'Managed web hosting with separate staging and production environments']]
    table = doc.add_table(rows=1, cols=2); table.style='Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(tech[0]): set_cell_text(table.rows[0].cells[i], h, 'FFFFFF', True, 9); shade(table.rows[0].cells[i], '1C4775')
    for row in tech[1:]:
        cells=table.add_row().cells
        for i,val in enumerate(row): set_cell_text(cells[i], val, size=9)
    add_heading(doc, '7. Success measures')
    for text in ['Public verification completed in under 30 seconds for a known certificate.', 'At least 95% of issued certificates discoverable through the public search.', 'Renewal reminders sent before an instrument enters the due-soon window.', 'Every status change attributable to a user, timestamp, and inspection event.', 'Reduced manual reconciliation between field inspections and department registers.']:
        p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3); p.add_run(text)
    add_heading(doc, '8. Delivery roadmap')
    roadmap = [['Phase', 'Deliverable'], ['1. Foundation', 'Data model, authentication, roles, certificate ID scheme'], ['2. Core portal', 'Registration, public lookup, certificate view, inspector queue'], ['3. Trust layer', 'Digital signatures, audit log, QR-ready certificate verification'], ['4. Pilot', 'One district rollout, field feedback, accessibility and security review'], ['5. Scale', 'District onboarding, analytics, renewal notifications, integrations']]
    table = doc.add_table(rows=1, cols=2); table.style='Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(roadmap[0]): set_cell_text(table.rows[0].cells[i], h, 'FFFFFF', True, 9); shade(table.rows[0].cells[i], '1C4775')
    for row in roadmap[1:]:
        cells=table.add_row().cells
        for i,val in enumerate(row): set_cell_text(cells[i], val, size=9)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('MeasureSure | SIH26036 | Prototype project brief').font.size = Pt(8)
    doc.save(DOCX)

def footer(canvas, doc):
    canvas.saveState(); canvas.setStrokeColor(colors.HexColor('#E5E9EE')); canvas.line(18*mm, 15*mm, 192*mm, 15*mm); canvas.setFont('Helvetica', 7.5); canvas.setFillColor(colors.HexColor('#8B96A3')); canvas.drawString(18*mm, 9*mm, 'MeasureSure | SIH26036'); canvas.drawRightString(192*mm, 9*mm, f'Page {doc.page}'); canvas.restoreState()

def build_pdf():
    doc = SimpleDocTemplate(str(PDF), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=17*mm, bottomMargin=20*mm)
    s = getSampleStyleSheet(); blue = colors.HexColor('#2563EB'); navy = colors.HexColor('#1C4775'); muted = colors.HexColor('#4D5B6C')
    title = ParagraphStyle('TitleX', parent=s['Title'], fontName='Helvetica-Bold', fontSize=27, leading=31, textColor=blue, alignment=TA_CENTER, spaceAfter=6)
    sub = ParagraphStyle('SubX', parent=s['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=15, textColor=navy, alignment=TA_CENTER, spaceAfter=7)
    h1 = ParagraphStyle('H1X', parent=s['Heading1'], fontName='Helvetica-Bold', fontSize=15, leading=18, textColor=navy, spaceBefore=13, spaceAfter=6)
    body = ParagraphStyle('BodyX', parent=s['BodyText'], fontName='Helvetica', fontSize=9.6, leading=14, textColor=muted, spaceAfter=7)
    th = ParagraphStyle('THX', parent=body, fontName='Helvetica-Bold', textColor=colors.white, spaceAfter=0)
    bullet = ParagraphStyle('BulletX', parent=body, leftIndent=12, firstLineIndent=-8, bulletIndent=0, spaceAfter=3)
    story = [Spacer(1, 16), Paragraph('MeasureSure', title), Paragraph('SIH26036 | Online Verification System for Weighing and Measuring Instruments', sub), Paragraph('Project brief and implementation concept | 29 September 2026', ParagraphStyle('date', parent=body, alignment=TA_CENTER, fontSize=8, textColor=colors.HexColor('#7D8896'))), Spacer(1, 12), Paragraph('MeasureSure is a secure, role-based digital portal that helps citizens, businesses, and Legal Metrology inspectors register, verify, monitor, and renew weighing and measuring instruments. The product is designed around a simple public promise: every measurement should be traceable to a current, trusted verification record.', body), Paragraph('1. Problem statement', h1), Paragraph('Verification records are often spread across paper certificates, local registers, and disconnected spreadsheets. This makes it difficult to confirm whether an instrument is valid at the point of use, creates avoidable work for inspectors, and increases the risk of expired or tampered certificates being accepted.', body), Paragraph('2. Proposed solution', h1), Paragraph('MeasureSure creates one digital record per instrument. Each record carries a unique certificate ID, owner and location details, instrument type, verification history, renewal date, and a digitally signed status. The public search experience is intentionally lightweight: a user can search by certificate number, serial number, or owner name and immediately see whether the instrument is valid, due soon, or pending review.', body)]
    data = [['User','Primary journey','Outcome'],['Citizen / buyer','Search certificate before purchase or transaction','Quick confidence in the reading'],['Business owner','Register instrument and track renewal due dates','Fewer missed renewals'],['Inspector','Review queue, verify on site, issue certificate','Faster, auditable field work'],['Department admin','Monitor district activity and exceptions','Clear operational visibility']]
    t = Table([[Paragraph(x, th) for x in data[0]]] + [[Paragraph(x,body) for x in r] for r in data[1:]], colWidths=[37*mm,86*mm,46*mm]); t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),navy),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#D9E1EA')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)])); story += [Paragraph('3. Primary users and journeys', h1), t, Paragraph('4. Product scope', h1)]
    for x in ['Public verification search with certificate, serial, and owner lookup.','Instrument registration with owner, location, type, and supporting evidence.','Inspector workspace with inspection queue, verification actions, and certificate issuance.','Certificate repository with status, audit trail, renewal dates, and downloadable records.','Dashboard metrics for registered, verified, due-soon, and pending instruments.','Tamper-evident digital signatures and role-based access control.']: story.append(Paragraph('• '+x, bullet))
    story += [Paragraph('5. Prototype interface', h1), Paragraph('The delivered website prototype uses a calm government-service visual language: a clear left navigation, a public verification panel, a compact operational dashboard, and a recent-records table. The main interaction is live in the browser: type a certificate or owner query to filter the record set, use the navigation to change workspace context, or trigger registration and record actions.', body), Paragraph('6. Suggested technical architecture', h1)]
    tech = [['Layer','Recommended implementation'],['Frontend','Next.js App Router, responsive React components, accessible form controls'],['API','REST or typed RPC endpoints for registration, lookup, verification, and renewal'],['Data','PostgreSQL with instrument, owner, certificate, inspection, and audit tables'],['Security','Role-based auth, signed certificate payloads, encrypted storage, immutable audit events'],['Deployment','Managed web hosting with separate staging and production environments']]
    t2 = Table([[Paragraph(x, th) for x in tech[0]]] + [[Paragraph(x,body) for x in r] for r in tech[1:]], colWidths=[34*mm,135*mm]); t2.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),navy),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#D9E1EA')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)])); story += [t2, Paragraph('7. Success measures', h1)]
    for x in ['Public verification completed in under 30 seconds for a known certificate.','At least 95% of issued certificates discoverable through the public search.','Renewal reminders sent before an instrument enters the due-soon window.','Every status change attributable to a user, timestamp, and inspection event.','Reduced manual reconciliation between field inspections and department registers.']: story.append(Paragraph('• '+x, bullet))
    story += [Paragraph('8. Delivery roadmap', h1)]
    roadmap = [['Phase','Deliverable'],['1. Foundation','Data model, authentication, roles, certificate ID scheme'],['2. Core portal','Registration, public lookup, certificate view, inspector queue'],['3. Trust layer','Digital signatures, audit log, QR-ready certificate verification'],['4. Pilot','One district rollout, field feedback, accessibility and security review'],['5. Scale','District onboarding, analytics, renewal notifications, integrations']]
    t3 = Table([[Paragraph(x, th) for x in roadmap[0]]] + [[Paragraph(x,body) for x in r] for r in roadmap[1:]], colWidths=[34*mm,135*mm]); t3.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),navy),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#D9E1EA')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)])); story.append(t3)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)

if __name__ == '__main__': build_docx(); build_pdf(); print(DOCX); print(PDF)
